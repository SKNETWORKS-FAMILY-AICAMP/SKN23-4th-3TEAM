"""
api/views/chat_views.py
기존 back/routers/chat_router.py → Django view 변환
SSE 스트리밍 + 단계별 진행 메시지 + 피부 상식 모달 연동 포함
"""
import os
import sys
import json
import random
import queue
import threading
import openpyxl
import requests as _requests
from django.http import JsonResponse, StreamingHttpResponse
from django.views.decorators.csrf import csrf_exempt
from api.middleware import require_auth, parse_json_body
from services import chat_service
from db.schemas import ChatRoomCreate, MessageCreate

# 프로젝트 루트를 sys.path에 추가
_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)

from ai.orchestrator.graph import run

# ─────────────────────────────────────────────
# 페르소나 문구 로드
# ─────────────────────────────────────────────

def _load_persona_messages():
    _BASE = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
    wb = openpyxl.load_workbook(os.path.join(_BASE, "assets", "ongle_contents_with_emoji.xlsx"))
    return {
        "loading":  [row[1] for row in wb["로딩멘트"].iter_rows(min_row=2, values_only=True)],
        "tips":     [row[1] for row in wb["피부정보꿀팁"].iter_rows(min_row=2, values_only=True)],
        "positive": [row[1] for row in wb["긍정메시지"].iter_rows(min_row=2, values_only=True)],
    }

PERSONA_MESSAGES = _load_persona_messages()

# ─────────────────────────────────────────────
# 내부 헬퍼
# ─────────────────────────────────────────────

def _msg_to_response(msg) -> dict:
    return {
        "message_id": msg.message_id,
        "chat_room_id": msg.chat_room_id,
        "role": msg.role,
        "model_type": msg.model_type,
        "content": msg.content,
        "created_at": msg.created_at.isoformat() if msg.created_at else None,
        "image_url": msg.image_urls or [],
    }

def _room_to_response(room) -> dict:
    latest = chat_service.get_latest_message_by_room(room.chat_room_id)
    return {
        "chat_room_id": room.chat_room_id,
        "user_id": room.user_id,
        "title": room.title,
        "created_at": room.created_at.isoformat() if room.created_at else None,
        "last_message": latest.content[:50] if latest and latest.content else None,
        "last_message_at": latest.created_at.isoformat() if latest and latest.created_at else None,
    }


def _run_ai(user_text, image_urls, model_type, user_id, chat_history, is_first_message, step_callback=None):
    _type_map = {
        "simple": "quick",
        "detailed": "detailed",
        "ingredient": "ingredient",
        "personal": "personal",
    }
    analysis_type = _type_map.get(model_type)

    if step_callback:
        step_callback("🔍 피부 고민을 분석하고 있어요...")

    image_bytes = []
    if image_urls and analysis_type in ("quick", "detailed", "ingredient", "personal"):
        if step_callback:
            step_callback("📸 업로드된 이미지를 확인하고 있어요...")
        for url in image_urls:
            try:
                resp = _requests.get(url, timeout=10)
                resp.raise_for_status()
                image_bytes.append(resp.content)
            except Exception as e:
                print(f"[chat_views] 이미지 다운로드 실패: {url} → {repr(e)}", flush=True)

    if step_callback:
        if analysis_type in ("quick", "detailed"):
            step_callback("🔬 AI가 피부 상태를 분석하고 있어요...")
        elif analysis_type == "ingredient":
            step_callback("🏷️ 전성분을 추출하고 있어요...")
        elif analysis_type == "personal":
            step_callback("🎨 퍼스널컬러를 분석하고 있어요...")
        else:
            step_callback("📚 23만 건의 피부 데이터에서 근거를 찾고 있어요...")

    report = run(
        user_text=user_text,
        images=image_bytes,
        analysis_type=analysis_type,
        user_id=user_id,
        chat_history=chat_history,
        is_first_message=is_first_message,
        image_urls=image_urls,
    )
    return report.get("chat_answer") or "답변을 생성하지 못했어요. 다시 시도해주세요."


def _run_ai_guest(user_text, chat_history=None, step_callback=None):
    if step_callback:
        step_callback("🔍 피부 고민을 분석하고 있어요...")

    if step_callback:
        step_callback("📚 23만 건의 피부 데이터에서 근거를 찾고 있어요...")

    report = run(
        user_text=user_text,
        images=[],
        analysis_type=None,
        user_id=None,
        chat_history=chat_history or [],
        is_first_message=not bool(chat_history),
        image_urls=[],
    )
    return report.get("chat_answer") or "답변을 생성하지 못했어요. 다시 시도해주세요."


# ─────────────────────────────────────────────
# SSE 헬퍼
# ─────────────────────────────────────────────

def _sse_event(data: dict) -> str:
    return f"data: {json.dumps(data, ensure_ascii=False)}\n\n"


# ─────────────────────────────────────────────
# Views
# ─────────────────────────────────────────────

@csrf_exempt
@require_auth
def chat_rooms(request, user_id):
    """POST: 채팅방 생성, GET: 채팅방 목록"""
    if request.method == "POST":
        data = ChatRoomCreate(user_id=user_id)
        room = chat_service.create_chat_room(data)
        return JsonResponse(_room_to_response(room), status=201)
    elif request.method == "GET":
        rooms = chat_service.get_chat_rooms_by_user(user_id)
        return JsonResponse([_room_to_response(r) for r in rooms], safe=False)
    return JsonResponse({"detail": "허용되지 않는 메서드입니다."}, status=405)


@csrf_exempt
def guest_message(request):
    """비로그인 텍스트 채팅"""
    if request.method != "POST":
        return JsonResponse({"detail": "허용되지 않는 메서드입니다."}, status=405)

    body = parse_json_body(request)
    if not body:
        return JsonResponse({"detail": "잘못된 요청입니다."}, status=400)

    user_text = (body.get("content") or "").strip()
    chat_history = body.get("chat_history") or []

    try:
        ai_text = _run_ai_guest(user_text, chat_history)
    except Exception as e:
        print(f"[guest_message AI ERROR] {repr(e)}", flush=True)
        ai_text = "잠시 후 다시 시도해주세요."

    persona = random.choice(PERSONA_MESSAGES.get("tips", ["피부 관리 팁을 확인해보세요!"]))
    return JsonResponse({
        "role": "assistant",
        "content": ai_text,
        "persona_tip": persona,
    })


@csrf_exempt
def guest_message_stream(request):
    """비로그인 SSE 스트리밍 채팅"""
    if request.method != "POST":
        return JsonResponse({"detail": "허용되지 않는 메서드입니다."}, status=405)

    body = parse_json_body(request)
    if not body:
        return JsonResponse({"detail": "잘못된 요청입니다."}, status=400)

    user_text = (body.get("content") or "").strip()
    chat_history = body.get("chat_history") or []

    def event_generator():
        step_queue = queue.Queue()

        def step_callback(msg):
            step_queue.put(msg)

        # 로딩 시작 신호
        yield _sse_event({"type": "loading", "message": random.choice(PERSONA_MESSAGES["loading"])})

        result_holder = {"text": None, "error": None}

        def run_ai_thread():
            try:
                result_holder["text"] = _run_ai_guest(user_text, chat_history, step_callback=step_callback)
            except Exception as e:
                result_holder["error"] = e
            finally:
                step_queue.put(None)

        thread = threading.Thread(target=run_ai_thread, daemon=True)
        thread.start()

        # 단계별 메시지 전송
        import time
        while True:
            try:
                msg = step_queue.get(timeout=0.3)
                if msg is None:
                    break
                yield _sse_event({"type": "loading", "message": msg})
            except queue.Empty:
                time.sleep(0.1)

        thread.join(timeout=60)

        try:
            if result_holder["error"]:
                raise result_holder["error"]
            yield _sse_event({"type": "done", "content": result_holder["text"]})
        except Exception as e:
            print(f"[guest_stream ERROR] {repr(e)}", flush=True)
            yield _sse_event({"type": "error", "message": "잠시 후 다시 시도해주세요."})

    response = StreamingHttpResponse(event_generator(), content_type="text/event-stream")
    response["Cache-Control"] = "no-cache"
    response["X-Accel-Buffering"] = "no"
    return response


@csrf_exempt
@require_auth
def chat_room_detail(request, chat_room_id, user_id):
    """GET: 채팅방 조회, DELETE: 채팅방 삭제"""
    room = chat_service.get_chat_room_by_id(chat_room_id)
    if not room:
        return JsonResponse({"detail": "채팅방을 찾을 수 없습니다."}, status=404)
    if room.user_id != user_id:
        return JsonResponse({"detail": "접근 권한이 없습니다."}, status=403)

    if request.method == "GET":
        return JsonResponse(_room_to_response(room))
    elif request.method == "DELETE":
        chat_service.delete_chat_room(chat_room_id)
        return JsonResponse({}, status=204)
    return JsonResponse({"detail": "허용되지 않는 메서드입니다."}, status=405)


@csrf_exempt
@require_auth
def chat_messages(request, chat_room_id, user_id):
    """POST: 메시지 전송, GET: 히스토리 조회"""
    room = chat_service.get_chat_room_by_id(chat_room_id)
    if not room:
        return JsonResponse({"detail": "채팅방을 찾을 수 없습니다."}, status=404)
    if room.user_id != user_id:
        return JsonResponse({"detail": "접근 권한이 없습니다."}, status=403)

    if request.method == "GET":
        role = request.GET.get("role")
        if role:
            msgs = chat_service.get_messages_by_role(chat_room_id, role)
        else:
            msgs = chat_service.get_messages_by_room(chat_room_id)
        return JsonResponse([_msg_to_response(m) for m in msgs], safe=False)

    elif request.method == "POST":
        raw = parse_json_body(request)
        if not raw:
            return JsonResponse({"detail": "잘못된 요청입니다."}, status=400)

        body = MessageCreate(
            chat_room_id=chat_room_id,
            role="user",
            model_type=raw.get("model_type", "default"),
            content=raw.get("content"),
            image_url=raw.get("image_url"),
        )

        try:
            user_msg = chat_service.save_message(body)
        except ValueError as e:
            return JsonResponse({"detail": str(e)}, status=400)

        is_first = not room.title
        if is_first and body.content:
            title = body.content[:30] + ("..." if len(body.content) > 30 else "")
            chat_service.update_chat_room_title(chat_room_id, title)

        history_msgs = chat_service.get_messages_by_room(chat_room_id)
        chat_history = [
            {"role": m.role, "content": m.content or ""}
            for m in history_msgs[:-1]
            if m.role in ("user", "assistant") and m.content
        ]

        image_urls = body.image_url or []
        user_text = body.content or ""

        try:
            ai_text = _run_ai(
                user_text=user_text,
                image_urls=image_urls,
                model_type=body.model_type,
                user_id=user_id,
                chat_history=chat_history,
                is_first_message=is_first,
            )
        except Exception as e:
            print(f"[send_message AI ERROR] {repr(e)}", flush=True)
            ai_text = "잠시 후 다시 시도해주세요."

        ai_msg_data = MessageCreate(
            chat_room_id=chat_room_id,
            role="assistant",
            model_type=body.model_type,
            content=ai_text,
        )
        ai_msg = chat_service.save_message(ai_msg_data)

        return JsonResponse([_msg_to_response(user_msg), _msg_to_response(ai_msg)], safe=False, status=201)

    return JsonResponse({"detail": "허용되지 않는 메서드입니다."}, status=405)


@csrf_exempt
@require_auth
def chat_messages_stream(request, chat_room_id, user_id):
    """회원 SSE 스트리밍 메시지 전송"""
    if request.method != "POST":
        return JsonResponse({"detail": "허용되지 않는 메서드입니다."}, status=405)

    room = chat_service.get_chat_room_by_id(chat_room_id)
    if not room:
        return JsonResponse({"detail": "채팅방을 찾을 수 없습니다."}, status=404)
    if room.user_id != user_id:
        return JsonResponse({"detail": "접근 권한이 없습니다."}, status=403)

    raw = parse_json_body(request)
    if not raw:
        return JsonResponse({"detail": "잘못된 요청입니다."}, status=400)

    body = MessageCreate(
        chat_room_id=chat_room_id,
        role="user",
        model_type=raw.get("model_type", "default"),
        content=raw.get("content"),
        image_url=raw.get("image_url"),
    )

    # 사용자 메시지 DB 저장
    try:
        user_msg = chat_service.save_message(body)
    except ValueError as e:
        return JsonResponse({"detail": str(e)}, status=400)

    # 첫 메시지 제목 설정
    is_first = not room.title
    if is_first and body.content:
        title = body.content[:30] + ("..." if len(body.content) > 30 else "")
        chat_service.update_chat_room_title(chat_room_id, title)

    # 히스토리 조회
    history_msgs = chat_service.get_messages_by_room(chat_room_id)
    chat_history = [
        {"role": m.role, "content": m.content or ""}
        for m in history_msgs[:-1]
        if m.role in ("user", "assistant") and m.content
    ]

    def event_generator():
        step_queue = queue.Queue()

        def step_callback(msg):
            step_queue.put(msg)

        # 로딩 시작 신호
        yield _sse_event({"type": "loading", "message": random.choice(PERSONA_MESSAGES["loading"])})

        result_holder = {"text": None, "error": None}

        def run_ai_thread():
            try:
                result_holder["text"] = _run_ai(
                    user_text=body.content or "",
                    image_urls=body.image_url or [],
                    model_type=body.model_type,
                    user_id=user_id,
                    chat_history=chat_history,
                    is_first_message=is_first,
                    step_callback=step_callback,
                )
            except Exception as e:
                result_holder["error"] = e
            finally:
                step_queue.put(None)

        thread = threading.Thread(target=run_ai_thread, daemon=True)
        thread.start()

        # 단계별 메시지 전송
        import time
        while True:
            try:
                msg = step_queue.get(timeout=0.3)
                if msg is None:
                    break
                yield _sse_event({"type": "loading", "message": msg})
            except queue.Empty:
                time.sleep(0.1)

        thread.join(timeout=60)

        try:
            if result_holder["error"]:
                raise result_holder["error"]

            ai_text = result_holder["text"]

            # AI 응답 DB 저장
            ai_msg = chat_service.save_message(MessageCreate(
                chat_room_id=chat_room_id,
                role="assistant",
                model_type=body.model_type,
                content=ai_text,
            ))

            yield _sse_event({"type": "done", "content": ai_text})

        except Exception as e:
            print(f"[stream ERROR] {repr(e)}", flush=True)
            yield _sse_event({"type": "error", "message": "잠시 후 다시 시도해주세요."})

    response = StreamingHttpResponse(event_generator(), content_type="text/event-stream")
    response["Cache-Control"] = "no-cache"
    response["X-Accel-Buffering"] = "no"
    return response
